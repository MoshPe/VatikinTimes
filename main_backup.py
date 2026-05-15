import sys
import os
import subprocess
import requests
import openpyxl
from datetime import datetime, timedelta
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from PyQt6.QtWidgets import (QApplication, QWidget, QVBoxLayout, QHBoxLayout,
                             QLabel, QComboBox, QTextEdit, QPushButton, QMessageBox)
from PyQt6.QtCore import Qt


class ZmanimApp(QWidget):
    def __init__(self):
        super().__init__()

        self.setWindowTitle("מחולל לוח זמני תפילה שבועי")
        self.resize(450, 550)

        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)

        layout = QVBoxLayout()
        layout.setSpacing(15)

        # 1. City Selection
        city_layout = QHBoxLayout()
        city_label = QLabel("בחר עיר:")
        city_label.setStyleSheet("font-weight: bold; font-size: 14px;")

        self.cities = {
            "קרית מוצקין": "293831",
            "קרית ים": "293822",
            "קרית אתא": "293845",
            "קרית ביאליק": "293844",
            "חיפה": "294801",
            "ירושלים": "281184",
            "תל אביב": "293397",
            "באר שבע": "295530"
        }
        self.city_dropdown = QComboBox()
        self.city_dropdown.addItems(list(self.cities.keys()))
        self.city_dropdown.setCurrentText("קרית מוצקין")

        city_layout.addWidget(city_label)
        city_layout.addWidget(self.city_dropdown)
        layout.addLayout(city_layout)

        # 2. Parasha Selection
        parasha_layout = QHBoxLayout()
        parasha_label = QLabel("בחר פרשת שבוע:")
        parasha_label.setStyleSheet("font-weight: bold; font-size: 14px;")

        self.parashiyot = {}
        self.parasha_dropdown = QComboBox()

        parasha_layout.addWidget(parasha_label)
        parasha_layout.addWidget(self.parasha_dropdown)
        layout.addLayout(parasha_layout)

        # 3. Comments Section
        comments_label = QLabel("הערות ללוח (יופיעו בתחתית):")
        comments_label.setStyleSheet("font-weight: bold; font-size: 14px;")
        layout.addWidget(comments_label)

        self.comments_text = QTextEdit()
        self.comments_text.setPlaceholderText("הקלד הערות כאן...")
        layout.addWidget(self.comments_text)

        # 4. Action Button
        self.generate_btn = QPushButton("עדכן והכן להדפסה")
        self.generate_btn.setStyleSheet("""
            QPushButton {
                background-color: #4CAF50;
                color: white;
                font-size: 16px;
                font-weight: bold;
                padding: 10px;
                border-radius: 5px;
            }
            QPushButton:disabled {
                background-color: #A5D6A7;
            }
        """)
        self.generate_btn.clicked.connect(self.generate_schedule)
        layout.addWidget(self.generate_btn)

        self.setLayout(layout)
        self.load_parashiyot()

    def get_upcoming_saturday(self):
        today = datetime.now()
        days_ahead = 5 - today.weekday()
        if days_ahead < 0:
            days_ahead += 7
        return (today + timedelta(days=days_ahead)).date()

    def load_parashiyot(self):
        try:
            current_year = datetime.now().year
            url = f"https://www.hebcal.com/hebcal?v=1&cfg=json&year={current_year}&s=on"
            response = requests.get(url)
            response.raise_for_status()
            data = response.json()

            today = datetime.now().date()
            upcoming_saturday = self.get_upcoming_saturday()

            for item in data.get("items", []):
                if item.get("category") == "parashat":
                    date_str = item.get("date")
                    parasha_date = datetime.strptime(date_str, "%Y-%m-%d").date()

                    if parasha_date >= today - timedelta(days=3):
                        heb_name = item.get("hebrew", item.get("title"))
                        self.parashiyot[heb_name] = date_str

            parasha_list = list(self.parashiyot.keys())
            self.parasha_dropdown.addItems(parasha_list)

            if parasha_list:
                default_index = 0
                for idx, name in enumerate(parasha_list):
                    if self.parashiyot[name] == upcoming_saturday.strftime("%Y-%m-%d"):
                        default_index = idx
                        break
                self.parasha_dropdown.setCurrentIndex(default_index)

        except Exception as e:
            QMessageBox.critical(self, "שגיאת רשת", f"לא ניתן למשוך רשימת פרשות.\n{e}")

    def fetch_week_zmanim(self, geoname_id, start_date_str):
        start_date = datetime.strptime(start_date_str, "%Y-%m-%d")

        zmanim_mapping = {
            "alotHaShachar": "עלות השחר",
            "misheyakir": "זמן טלית ותפילין",
            "sunrise": "הנץ החמה (הנראה)"
        }

        week_data = {hebrew_name: [] for hebrew_name in zmanim_mapping.values()}
        hebrew_dates = []

        try:
            for i in range(7):
                self.generate_btn.setText(f"מושך נתונים מ-Hebcal... (יום {i + 1}/7)")
                QApplication.processEvents()

                current_date = start_date + timedelta(days=i)
                date_str = current_date.strftime("%Y-%m-%d")

                # Added &sec=1 to the URL
                zmanim_url = f"https://www.hebcal.com/zmanim?cfg=json&geonameid={geoname_id}&date={date_str}&sec=1"
                zmanim_response = requests.get(zmanim_url)
                zmanim_response.raise_for_status()
                zmanim_data = zmanim_response.json()
                times = zmanim_data.get("times", {})

                for key, hebrew_name in zmanim_mapping.items():
                    if key in times:
                        time_iso = times[key]
                        time_str = time_iso.split('+')[0]
                        dt_obj = datetime.strptime(time_str, "%Y-%m-%dT%H:%M:%S")
                        # Updated to include seconds
                        week_data[hebrew_name].append(dt_obj.strftime("%H:%M:%S"))
                    else:
                        week_data[hebrew_name].append("--:--:--")  # Updated fallback

                converter_url = f"https://www.hebcal.com/converter?cfg=json&date={date_str}&g2h=1&strict=1"
                converter_response = requests.get(converter_url)
                converter_response.raise_for_status()
                converter_data = converter_response.json()

                hebrew_date_str = converter_data.get("hebrew", "")
                hebrew_dates.append(hebrew_date_str)

            return start_date, week_data, zmanim_mapping, hebrew_dates

        except Exception as e:
            QMessageBox.critical(self, "שגיאת רשת", f"שגיאה במשיכת הנתונים: {e}")
            return None, None, None, None

    def open_file_in_os(self, filename):
        try:
            if sys.platform == "win32":
                os.startfile(filename)
            elif sys.platform == "darwin":
                subprocess.call(["open", filename])
            else:
                subprocess.call(["xdg-open", filename])
        except Exception as e:
            QMessageBox.warning(self, "אזהרה", f"הקובץ נוצר אך לא ניתן היה לפתוח אותו אוטומטית.\n{e}")

    def generate_schedule(self):
        city_name = self.city_dropdown.currentText()
        geoname_id = self.cities[city_name]

        parasha_name = self.parasha_dropdown.currentText()
        if not parasha_name:
            QMessageBox.warning(self, "אזהרה", "אנא בחר פרשת שבוע תקינה.")
            return

        start_date_str = self.parashiyot[parasha_name]
        comments = self.comments_text.toPlainText().strip()

        self.generate_btn.setEnabled(False)

        start_date, week_data, zmanim_mapping, hebrew_dates = self.fetch_week_zmanim(geoname_id, start_date_str)

        if week_data and hebrew_dates:
            try:
                # Try to create and save the file
                filename = self.export_to_excel(city_name, parasha_name, start_date, week_data, comments,
                                                zmanim_mapping, hebrew_dates)
                self.open_file_in_os(filename)
                QMessageBox.information(self, "הצלחה", "הלוח השבועי עודכן ומוכן להדפסה!")

            except PermissionError:
                # Catch the Windows file lock error
                QMessageBox.critical(self, "שגיאת שמירה",
                                     "הקובץ כבר פתוח בתוכנה אחרת (כנראה Excel).\nאנא סגור את הקובץ הפתוח ונסה שוב.")
            except Exception as e:
                # Catch any other unexpected saving errors
                QMessageBox.critical(self, "שגיאה כללית", f"אירעה שגיאה בשמירת הקובץ:\n{e}")

        self.generate_btn.setText("עדכן והכן להדפסה")
        self.generate_btn.setEnabled(True)

    def export_to_excel(self, city, parasha_name, start_date, week_data, comments, zmanim_mapping, hebrew_dates):
        date_str = start_date.strftime("%Y-%m-%d")
        safe_parasha_name = "".join(c for c in parasha_name if c.isalnum() or c in " -")
        filename = f"לוח_זמנים_שבועי_{safe_parasha_name}_{date_str}.xlsx"
        clean_parasha = parasha_name.replace("פרשת ", "").strip()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "לוח זמנים"
        ws.sheet_view.rightToLeft = True

        # Header
        ws.append(["בס\"ד", "", "", "", "", "", "", "", date_str])
        ws.append([])

        # Main Title (Shifted right to column 2)
        ws.append(["", "בית הכנסת משמרת  --  ותיקין"])
        ws.cell(row=3, column=2).font = Font(size=20, bold=True)
        ws.append([])

        # Subtitle (Shifted right to column 3)
        ws.append(["", "", f'לוח זמנים לשבת פרשת "{clean_parasha}"'])
        ws.cell(row=5, column=3).font = Font(size=16, bold=True)
        ws.append([])

        # Row 7: Hebrew Dates (Reduced to 2 empty spaces so it starts at column 3 / C)
        ws.append(["", ""] + hebrew_dates)
        for col_idx in range(3, 10):
            cell = ws.cell(row=7, column=col_idx)
            cell.font = Font(size=11, bold=False, italic=True)
            cell.alignment = Alignment(horizontal="center")

        # Row 8: Days of the week
        ws.append(["", "", "שבת", "א'", "ב'", "ג'", "ד'", "ה'", "ו'"])
        for col_idx in range(3, 10):
            cell = ws.cell(row=8, column=col_idx)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        ws.append([])  # Row 9: Empty

        # Widen columns 3 through 9 (C through I)
        for col in range(3, 10):
            ws.column_dimensions[get_column_letter(col)].width = 16

        # Zmanim Data Rows
        current_row = 10
        for hebrew_name in zmanim_mapping.values():
            # Reduced to 1 empty space after the name to align with column 3
            row_data = [hebrew_name, ""] + week_data[hebrew_name]
            ws.append(row_data)

            for col_idx in range(3, 10):
                ws.cell(row=current_row, column=col_idx).alignment = Alignment(horizontal="center")

            ws.append([])
            current_row += 2

        # Calculate Shacharit Vatikin automatically
        vatikin_times = []
        netz_times = week_data.get("הנץ החמה (הנראה)", ["--:--"] * 7)

        for i, netz_str in enumerate(netz_times):
            if netz_str == "--:--":
                vatikin_times.append("--:--")
            else:
                netz_dt = datetime.strptime(netz_str, "%H:%M:%S")
                offset_minutes = 30 if i == 0 else 20
                vatikin_dt = netz_dt - timedelta(minutes=offset_minutes)
                vatikin_times.append(vatikin_dt.strftime("%H:%M:%S"))

        # Reduced to 1 empty space after the name
        ws.append(["שחרית ותיקין", ""] + vatikin_times)
        for col_idx in range(3, 10):
            ws.cell(row=current_row, column=col_idx).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=col_idx).font = Font(bold=True)

        ws.append([])

        ws.append(["", "הערות:"])
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True)

        if comments:
            for line in comments.split('\n'):
                ws.append(["", "", line])

        ws.append([])
        ws.append(["", "", "שבת שלום ומבורך בשורות טובות"])
        ws.cell(row=ws.max_row, column=3).font = Font(size=14, bold=True)

        wb.save(filename)
        return filename


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = ZmanimApp()
    window.show()
    sys.exit(app.exec())