import os
import sys
import types
import unittest
from datetime import datetime
from unittest import mock

import openpyxl

try:
    import PyQt6  # noqa: F401
except ModuleNotFoundError:
    qtwidgets = types.ModuleType("PyQt6.QtWidgets")
    qtcore = types.ModuleType("PyQt6.QtCore")

    class _DummyWidget:
        def __init__(self, *args, **kwargs):
            pass

        def __getattr__(self, name):
            return lambda *args, **kwargs: None

    class _DummyQt:
        class LayoutDirection:
            RightToLeft = None

    for name in (
        "QApplication",
        "QWidget",
        "QVBoxLayout",
        "QHBoxLayout",
        "QLabel",
        "QComboBox",
        "QTextEdit",
        "QSpinBox",
        "QPushButton",
        "QMessageBox",
    ):
        setattr(qtwidgets, name, _DummyWidget)
    qtcore.Qt = _DummyQt
    sys.modules["PyQt6"] = types.ModuleType("PyQt6")
    sys.modules["PyQt6.QtWidgets"] = qtwidgets
    sys.modules["PyQt6.QtCore"] = qtcore

import main


class SpecialEventsTest(unittest.TestCase):
    def test_special_events_are_aligned_to_week_days_and_ignore_learning_items(self):
        items = [
            {"date": "2026-05-15", "category": "holiday", "hebrew": "יום ירושלים"},
            {"date": "2026-05-15", "category": "dafyomi", "hebrew": "חולין דף ט״ו"},
            {"date": "2026-05-09", "category": "parashat", "hebrew": "פרשת בהר־בחקתי"},
            {
                "date": "2026-05-09",
                "category": "molad",
                "title": "מוֹלָד הָלְּבָנָה סִיוָן",
                "molad": {"dow": 6, "hour": 18, "minutes": 2, "chalakim": 15},
            },
            {
                "date": "2026-05-14",
                "category": "holiday",
                "hebrew": "ערב שבועות",
            },
            {
                "date": "2026-05-14T19:13:27+03:00",
                "category": "candles",
                "title_orig": "Candle lighting",
                "memo": "Erev Shavuot",
            },
            {
                "date": "2026-05-15T19:14:00+03:00",
                "category": "candles",
                "title_orig": "Candle lighting",
                "memo": "Parashat Nasso",
            },
        ]

        events = main.ZmanimApp.build_special_events_week(datetime(2026, 5, 9), items)

        self.assertEqual(
            [
                "המולד ליל ראשון 18:02 ו-15 חלקים",
                "--",
                "--",
                "--",
                "--",
                "ערב שבועות\nכניסת החג 19:13:27",
                "יום ירושלים",
            ],
            events,
        )

    def test_hebrew_year_is_split_from_daily_dates(self):
        dates = [
            "כ״ב בְּאִיָיר תשפ״ו",
            "כ״ג בְּאִיָיר תשפ״ו",
            "כ״ד בְּאִיָיר תשפ״ו",
        ]

        year, daily_dates = main.ZmanimApp.split_hebrew_year_from_dates(dates)

        self.assertEqual("תשפ״ו", year)
        self.assertEqual(["כ״ב בְּאִיָיר", "כ״ג בְּאִיָיר", "כ״ד בְּאִיָיר"], daily_dates)

    def test_yom_tov_week_marks_actual_yom_tov_but_not_erev_yom_tov(self):
        items = [
            {"date": "2026-05-21", "category": "holiday", "hebrew": "ערב שבועות"},
            {"date": "2026-05-22", "category": "holiday", "hebrew": "שבועות", "yomtov": True},
        ]

        yom_tov_week = main.ZmanimApp.build_yom_tov_week(datetime(2026, 5, 16), items)

        self.assertEqual([False, False, False, False, False, False, True], yom_tov_week)

    def test_hebcal_week_params_use_selected_town_for_candle_lighting(self):
        params = main.ZmanimApp.build_hebcal_week_params("281184", "2026-05-16", "2026-05-22")

        self.assertEqual("on", params["c"])
        self.assertEqual("geoname", params["geo"])
        self.assertEqual("281184", params["geonameid"])

    def test_parasha_list_params_use_israel_schedule(self):
        params = main.ZmanimApp.build_parasha_list_params(2026)

        self.assertEqual("2026", params["year"])
        self.assertEqual("on", params["s"])
        self.assertEqual("on", params["i"])
        self.assertEqual("he", params["lg"])

    def test_haftara_is_extracted_from_selected_shabbat_leyning(self):
        items = [
            {"date": "2026-05-21", "weekday": {"1": "Numbers 1:1-1:19"}},
            {
                "date": "2026-05-23",
                "haft": {"k": "I Samuel", "b": "20:18", "e": "20:42", "v": 25},
                "haftara": "I Samuel 20:18-42",
                "haftara_start": "ויאמר לו יהונתן",
            },
        ]

        haftara = main.ZmanimApp.extract_haftara_for_date(items, "2026-05-23")

        self.assertEqual("שמואל א׳ פרק כ׳ פסוקים י״ח-מ״ב - ויאמר לו יהונתן", haftara)

    def test_hebrew_opening_words_are_extracted_without_vowels_or_trope(self):
        verse = "וַיֹּאמֶר־ל֥וֹ יְהוֹנָתָ֖ן מָחָ֣ר חֹ֑דֶשׁ"

        opening_words = main.ZmanimApp.extract_opening_words(verse)

        self.assertEqual("ויאמר לו יהונתן", opening_words)

    def test_special_events_row_height_grows_for_multiline_content(self):
        height = main.ZmanimApp.calculate_content_row_height([
            "המולד ליל ראשון 18:02 ו-15 חלקים",
            "ערב שבועות\nכניסת החג 19:13:27",
            "--",
        ])

        self.assertGreaterEqual(height, 72)

    def test_week_zmanim_calculates_alot_and_talit_from_visible_netz_offsets(self):
        class DummyButton:
            def setText(self, text):
                pass

        class DummySpinBox:
            def __init__(self, value):
                self._value = value

            def value(self):
                return self._value

        class DummyResponse:
            def __init__(self, payload, status_code=200):
                self._payload = payload
                self.status_code = status_code

            def json(self):
                return self._payload

            def raise_for_status(self):
                pass

        def fake_get(url, params=None):
            if url == "https://www.hebcal.com/hebcal":
                return DummyResponse({"items": []})
            if url == "https://www.hebcal.com/leyning":
                return DummyResponse({"items": []})
            if url.startswith("https://www.hebcal.com/converter"):
                day = int(url.split("date=2026-05-")[1][:2]) - 15
                return DummyResponse({
                    "hebrew": f"יום {day}",
                    "hy": 5786,
                    "hm": "Iyyar",
                    "hd": day,
                })
            if url.startswith("https://www.hebcal.com/zmanim"):
                return DummyResponse({
                    "times": {
                        "alotHaShachar": "2026-05-16T03:00:00+03:00",
                        "misheyakir": "2026-05-16T04:00:00+03:00",
                        "sunrise": "2026-05-16T06:30:00+03:00",
                    }
                })
            raise AssertionError(f"Unexpected URL: {url}")

        app = main.ZmanimApp.__new__(main.ZmanimApp)
        app.generate_btn = DummyButton()
        app.chai_tables_cache = {
            5786: {
                "Iyyar": {day: "06:00:00" for day in range(1, 8)}
            }
        }
        app.alot_offset_spin = DummySpinBox(72)
        app.talit_tefillin_offset_spin = DummySpinBox(45)

        with mock.patch.object(main.requests, "get", side_effect=fake_get), \
                mock.patch.object(main.QApplication, "processEvents", return_value=None, create=True):
            _, week_data, _, _, _, _, _, _ = app.fetch_week_zmanim("294801", "חיפה", "2026-05-16")

        self.assertEqual(["04:48:00"] * 7, week_data["עלות השחר"])
        self.assertEqual(["05:15:00"] * 7, week_data["זמן טלית ותפילין"])
        self.assertEqual(["06:00:00"] * 7, week_data["הנץ החמה (הנראה)"])

    def test_export_writes_special_events_row_between_daf_yomi_and_comments(self):
        week_data = {
            "עלות השחר": ["04:23:54"] * 7,
            "זמן טלית ותפילין": ["04:49:35"] * 7,
            "הנץ החמה (הנראה)": ["05:51:05"] * 7,
        }
        zmanim_mapping = {
            "alotHaShachar": "עלות השחר",
            "misheyakir": "זמן טלית ותפילין",
            "sunrise": "הנץ החמה (הנראה)",
        }
        hebrew_dates = ["כ״ב בְּאִיָיר תשפ״ו"] * 7
        daf_yomi = ["חולין דף ט׳"] * 7
        special_events = [
            "המולד ליל ראשון 18:02 ו-15 חלקים",
            "ערב שבועות\nכניסת החג 19:13:27",
            "--",
            "--",
            "--",
            "--",
            "יום ירושלים",
        ]

        filename = "לוח_זמנים_שבועי_פרשת בהרבחקתי_2026-05-16.xlsx"
        if os.path.exists(filename):
            try:
                os.remove(filename)
            except PermissionError:
                pass

        try:
            filename = main.ZmanimApp.export_to_excel(
                None,
                "קריות",
                "פרשת בהר־בחקתי",
                datetime(2026, 5, 16),
                week_data,
                "",
                zmanim_mapping,
                hebrew_dates,
                daf_yomi,
                special_events,
                "שמואל א׳ פרק כ׳ פסוקים י״ח-מ״ב - ויאמר לו יהונתן",
                [False, False, True, False, False, False, False],
            )
            wb = openpyxl.load_workbook(filename)
            ws = wb.active
            values = {
                "B2": ws["B2"].value,
                "B5": ws["B5"].value,
                "C7": ws["C7"].value,
                "D16": ws["D16"].value,
                "E16": ws["E16"].value,
                "A18": ws["A18"].value,
                "A20": ws["A20"].value,
                "I20": ws["I20"].value,
                "B22": ws["B22"].value,
            }
            styles = {
                "B5_font_name": ws["B5"].font.name,
                "B5_underline": ws["B5"].font.underline,
                "C7_font_name": ws["C7"].font.name,
                "C7_font_size": ws["C7"].font.sz,
                "C8_font_name": ws["C8"].font.name,
                "A20_wrap_text": ws["A20"].alignment.wrap_text,
                "A20_shrink_to_fit": ws["A20"].alignment.shrink_to_fit,
                "I20_wrap_text": ws["I20"].alignment.wrap_text,
                "I20_shrink_to_fit": ws["I20"].alignment.shrink_to_fit,
                "A20_row_height": ws.row_dimensions[20].height,
                "A14_row_height": ws.row_dimensions[14].height,
                "A14_wrap_text": ws["A14"].alignment.wrap_text,
                "A_column_width": ws.column_dimensions["A"].width,
                "B_column_width": ws.column_dimensions["B"].width,
                "C_column_width": ws.column_dimensions["C"].width,
                "fit_to_page": ws.sheet_properties.pageSetUpPr.fitToPage,
                "fit_to_width": ws.page_setup.fitToWidth,
                "fit_to_height": ws.page_setup.fitToHeight,
                "print_area": ws.print_area,
            }
            merged_ranges = {str(merged_range) for merged_range in ws.merged_cells.ranges}
            wb.close()
        finally:
            if os.path.exists(filename):
                try:
                    os.remove(filename)
                except PermissionError:
                    pass

        self.assertEqual("תשפ״ו", values["B2"])
        self.assertEqual('לוח זמנים לשבת פרשת "בהר־בחקתי" | הפטרה: שמואל א׳ פרק כ׳ פסוקים י״ח-מ״ב - ויאמר לו יהונתן', values["B5"])
        self.assertEqual("כ״ב בְּאִיָיר", values["C7"])
        self.assertEqual("05:29:05", values["D16"])
        self.assertEqual("05:19:05", values["E16"])
        self.assertEqual("דף יומי", values["A18"])
        self.assertEqual("ימים מיוחדים", values["A20"])
        self.assertEqual("יום ירושלים", values["I20"])
        self.assertEqual("הערות:", values["B22"])
        self.assertEqual("David", styles["B5_font_name"])
        self.assertEqual("single", styles["B5_underline"])
        self.assertEqual("David", styles["C7_font_name"])
        self.assertEqual(14, styles["C7_font_size"])
        self.assertEqual("David", styles["C8_font_name"])
        self.assertTrue(styles["A20_wrap_text"])
        self.assertFalse(styles["A20_shrink_to_fit"])
        self.assertTrue(styles["I20_wrap_text"])
        self.assertFalse(styles["I20_shrink_to_fit"])
        self.assertTrue(styles["A14_wrap_text"])
        self.assertGreaterEqual(styles["A14_row_height"], 36)
        self.assertGreaterEqual(styles["A20_row_height"], 72)
        self.assertGreaterEqual(styles["A_column_width"], 18)
        self.assertGreaterEqual(styles["B_column_width"], 3)
        self.assertGreaterEqual(styles["C_column_width"], 17)
        self.assertTrue(styles["fit_to_page"])
        self.assertEqual(1, styles["fit_to_width"])
        self.assertEqual(1, styles["fit_to_height"])
        self.assertEqual(["$A$1:$I$24"], styles["print_area"])
        self.assertIn("B2:I2", merged_ranges)
        self.assertIn("B3:I3", merged_ranges)
        self.assertIn("B5:I5", merged_ranges)


if __name__ == "__main__":
    unittest.main()
