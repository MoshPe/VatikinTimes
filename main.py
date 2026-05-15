import sys
import os
import subprocess
import requests
import unicodedata
from bs4 import BeautifulSoup
import openpyxl
from datetime import datetime, timedelta
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from PyQt6.QtWidgets import (QApplication, QWidget, QVBoxLayout, QHBoxLayout,
                             QLabel, QComboBox, QTextEdit, QPushButton, QMessageBox)
from PyQt6.QtCore import Qt


class ZmanimApp(QWidget):
    HEBREW_BOOK_NAMES = {
        "Genesis": "בראשית",
        "Exodus": "שמות",
        "Leviticus": "ויקרא",
        "Numbers": "במדבר",
        "Deuteronomy": "דברים",
        "Joshua": "יהושע",
        "Judges": "שופטים",
        "I Samuel": "שמואל א׳",
        "II Samuel": "שמואל ב׳",
        "I Kings": "מלכים א׳",
        "II Kings": "מלכים ב׳",
        "Isaiah": "ישעיהו",
        "Jeremiah": "ירמיהו",
        "Ezekiel": "יחזקאל",
        "Hosea": "הושע",
        "Joel": "יואל",
        "Amos": "עמוס",
        "Obadiah": "עובדיה",
        "Jonah": "יונה",
        "Micah": "מיכה",
        "Nahum": "נחום",
        "Habakkuk": "חבקוק",
        "Zephaniah": "צפניה",
        "Haggai": "חגי",
        "Zechariah": "זכריה",
        "Malachi": "מלאכי",
    }
    HEBREW_HAFTARA_OPENINGS = {
        ("I Samuel", "20:18"): "ויאמר לו יהונתן",
    }

    def __init__(self):
        super().__init__()

        self.setWindowTitle("מחולל לוח זמני תפילה שבועי - כולל דף יומי")
        self.resize(450, 550)
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)

        layout = QVBoxLayout()
        layout.setSpacing(15)

        # 1. City Selection (Added Krayot based on the Chai Tables payload)
        city_layout = QHBoxLayout()
        city_label = QLabel("בחר עיר:")
        city_label.setStyleSheet("font-weight: bold; font-size: 14px;")

        # Dictionary maps to (Hebcal_ID, ChaiTables_MetroArea)
        self.cities = {
            "קריות": ("294801", "קרית_ים-מוצקין-ביאליק"),  # Using Haifa coordinates for Hebcal fallback
            "חיפה": ("294801", "חיפה"),
            "ירושלים": ("281184", "ירושלים"),
            "תל אביב": ("293397", "תל_אביב-יפו"),
            "באר שבע": ("295530", "באר_שבע")
        }
        self.city_dropdown = QComboBox()
        self.city_dropdown.addItems(list(self.cities.keys()))
        self.city_dropdown.setCurrentText("קריות")

        city_layout.addWidget(city_label)
        city_layout.addWidget(self.city_dropdown)
        layout.addLayout(city_layout)

        # 2. Parasha Selection
        parasha_layout = QHBoxLayout()
        parasha_label = QLabel("בחר פרשת שבוע:")
        parasha_label.setStyleSheet("font-weight: bold; font-size: 14px;")

        self.parashiyot = {}
        self.parasha_dropdown = QComboBox()

        parasha_layout.addWidget(parasha_label)
        parasha_layout.addWidget(self.parasha_dropdown)
        layout.addLayout(parasha_layout)

        # 3. Comments Section
        comments_label = QLabel("הערות ללוח (יופיעו בתחתית):")
        comments_label.setStyleSheet("font-weight: bold; font-size: 14px;")
        layout.addWidget(comments_label)

        self.comments_text = QTextEdit()
        self.comments_text.setPlaceholderText("הקלד הערות כאן...")
        layout.addWidget(self.comments_text)

        # 4. Action Button
        self.generate_btn = QPushButton("עדכן והכן להדפסה")
        self.generate_btn.setStyleSheet("""
            QPushButton {
                background-color: #4CAF50;
                color: white;
                font-size: 16px;
                font-weight: bold;
                padding: 10px;
                border-radius: 5px;
            }
            QPushButton:disabled {
                background-color: #A5D6A7;
            }
        """)
        self.generate_btn.clicked.connect(self.generate_schedule)
        layout.addWidget(self.generate_btn)

        self.setLayout(layout)

        # Cache for Chai Tables
        self.chai_tables_cache = {}

        self.load_parashiyot()

    @staticmethod
    def build_special_events_week(start_date, items):
        if isinstance(start_date, datetime):
            start_date = start_date.date()

        allowed_categories = {"holiday", "roshchodesh"}
        events_by_date = {}
        holiday_dates = {
            item.get("date", "")[:10]
            for item in items
            if item.get("category") == "holiday"
            and (item.get("hebrew", "").startswith("ערב ") or item.get("title", "").startswith("Erev "))
        }

        for item in items:
            event_date = item.get("date", "")[:10]
            if not event_date:
                continue

            event_name = None
            category = item.get("category")

            if category in allowed_categories:
                event_name = item.get("hebrew") or item.get("title")
            elif category in {"molad", "mevarchim"}:
                event_name = ZmanimApp.extract_molad_text(item)
            elif category == "candles" and event_date in holiday_dates:
                event_name = ZmanimApp.format_yom_tov_candle_lighting(item)

            if not event_name:
                continue

            events_by_date.setdefault(event_date, []).append(event_name)

        special_events = []
        for i in range(7):
            date_str = (start_date + timedelta(days=i)).strftime("%Y-%m-%d")
            special_events.append("\n".join(events_by_date.get(date_str, [])) or "--")

        return special_events

    @staticmethod
    def extract_molad_text(item):
        molad = item.get("molad")
        if molad:
            return ZmanimApp.format_molad_text(molad)

        for field in ("hebrew", "memo", "title"):
            value = item.get(field, "")
            if "המולד" in value:
                return value
        return ""

    @staticmethod
    def format_molad_text(molad):
        dow = int(molad.get("dow", 0))
        hour = int(molad.get("hour", 0))
        minutes = int(molad.get("minutes", 0))
        chalakim = int(molad.get("chalakim", 0))

        day_names = {
            0: "ראשון",
            1: "שני",
            2: "שלישי",
            3: "רביעי",
            4: "חמישי",
            5: "שישי",
            6: "שבת",
        }

        if hour >= 18:
            day_text = f"ליל {day_names[(dow + 1) % 7]}"
        else:
            day_text = f"יום {day_names[dow]}"

        return f"המולד {day_text} {hour:02d}:{minutes:02d} ו-{chalakim} חלקים"

    @staticmethod
    def format_yom_tov_candle_lighting(item):
        date_value = item.get("date", "")
        try:
            candle_dt = datetime.fromisoformat(date_value)
        except ValueError:
            return ""

        candle_time = candle_dt.strftime("%H:%M:%S" if candle_dt.second else "%H:%M")
        return f"כניסת החג {candle_time}"

    @staticmethod
    def calculate_content_row_height(values, column_width=17):
        max_lines = 1
        for value in values:
            if not value or value == "--":
                continue

            value_lines = 0
            for line in str(value).splitlines() or [""]:
                estimated_lines = max(1, (len(line) + column_width - 1) // column_width)
                value_lines += estimated_lines
            max_lines = max(max_lines, value_lines)

        return max(30, max_lines * 30)

    @staticmethod
    def build_yom_tov_week(start_date, items):
        if isinstance(start_date, datetime):
            start_date = start_date.date()

        yom_tov_dates = {
            item.get("date", "")[:10]
            for item in items
            if item.get("yomtov") is True
        }

        return [
            (start_date + timedelta(days=i)).strftime("%Y-%m-%d") in yom_tov_dates
            for i in range(7)
        ]

    @staticmethod
    def build_hebcal_week_params(hebcal_id, start_date_str, end_date_str):
        return {
            "cfg": "json",
            "v": "1",
            "F": "on",
            "maj": "on",
            "min": "on",
            "mf": "on",
            "nx": "on",
            "ss": "on",
            "mod": "on",
            "molad": "on",
            "c": "on",
            "geo": "geoname",
            "geonameid": hebcal_id,
            "i": "on",
            "lg": "he",
            "start": start_date_str,
            "end": end_date_str,
        }

    @staticmethod
    def build_parasha_list_params(year):
        return {
            "v": "1",
            "cfg": "json",
            "year": str(year),
            "s": "on",
            "i": "on",
            "lg": "he",
        }

    @staticmethod
    def split_hebrew_year_from_dates(hebrew_dates):
        if not hebrew_dates:
            return "", []

        split_dates = []
        years = []
        for hebrew_date in hebrew_dates:
            parts = hebrew_date.rsplit(maxsplit=1)
            if len(parts) == 2:
                split_dates.append(parts[0])
                years.append(parts[1])
            else:
                split_dates.append(hebrew_date)

        if years and len(years) == len(hebrew_dates) and len(set(years)) == 1:
            return years[0], split_dates

        return "", hebrew_dates

    @staticmethod
    def format_hebrew_number(number):
        number = int(number)
        if number <= 0:
            return str(number)

        ones = {
            1: "א", 2: "ב", 3: "ג", 4: "ד", 5: "ה",
            6: "ו", 7: "ז", 8: "ח", 9: "ט",
        }
        tens = {
            10: "י", 20: "כ", 30: "ל", 40: "מ", 50: "נ",
            60: "ס", 70: "ע", 80: "פ", 90: "צ",
        }
        hundreds = {
            100: "ק", 200: "ר", 300: "ש", 400: "ת",
        }

        letters = []
        while number >= 400:
            letters.append(hundreds[400])
            number -= 400

        for value in (300, 200, 100):
            if number >= value:
                letters.append(hundreds[value])
                number -= value

        if number == 15:
            letters.extend(["ט", "ו"])
        elif number == 16:
            letters.extend(["ט", "ז"])
        else:
            for value in (90, 80, 70, 60, 50, 40, 30, 20, 10):
                if number >= value:
                    letters.append(tens[value])
                    number -= value
                    break
            if number:
                letters.append(ones[number])

        if len(letters) == 1:
            return f"{letters[0]}׳"

        return "".join(letters[:-1]) + "״" + letters[-1]

    @classmethod
    def format_hebrew_tanach_ref(cls, book, start_ref, end_ref):
        hebrew_book = cls.HEBREW_BOOK_NAMES.get(book, book)
        start_chapter, start_verse = start_ref.split(":", 1)
        end_chapter, end_verse = end_ref.split(":", 1)

        start_chapter_he = cls.format_hebrew_number(start_chapter)
        start_verse_he = cls.format_hebrew_number(start_verse)
        end_chapter_he = cls.format_hebrew_number(end_chapter)
        end_verse_he = cls.format_hebrew_number(end_verse)

        if start_chapter == end_chapter:
            ref = f"פרק {start_chapter_he} פסוקים {start_verse_he}-{end_verse_he}"
        else:
            ref = f"פרק {start_chapter_he} פסוק {start_verse_he} - פרק {end_chapter_he} פסוק {end_verse_he}"

        return f"{hebrew_book} {ref}"

    @staticmethod
    def extract_opening_words(hebrew_text, word_count=3):
        text = BeautifulSoup(hebrew_text or "", "html.parser").get_text(" ")
        text = text.replace("־", " ")
        text = "".join(
            char for char in unicodedata.normalize("NFD", text)
            if unicodedata.category(char) != "Mn"
        )
        words = [word.strip(".,;:!?()[]{}\"׳״") for word in text.split()]
        words = [word for word in words if word]
        return " ".join(words[:word_count])

    @staticmethod
    def build_sefaria_ref(book, start_ref):
        return f"{book.replace(' ', '_')}.{start_ref.replace(':', '.')}"

    @classmethod
    def fetch_haftara_opening_words(cls, haft):
        sefaria_ref = cls.build_sefaria_ref(haft["k"], haft["b"])
        response = requests.get(
            f"https://www.sefaria.org/api/texts/{sefaria_ref}",
            params={"context": 0},
        )
        response.raise_for_status()
        return cls.extract_opening_words(response.json().get("he", ""))

    @classmethod
    def extract_haftara_opening_for_item(cls, item, include_opening_lookup=False):
        if item.get("haftara_start"):
            return item["haftara_start"]

        haft = item.get("haft", {})
        if not (haft.get("k") and haft.get("b")):
            return ""

        fallback = cls.HEBREW_HAFTARA_OPENINGS.get((haft["k"], haft["b"]))
        if fallback:
            return fallback

        if include_opening_lookup:
            try:
                return cls.fetch_haftara_opening_words(haft)
            except (requests.RequestException, ValueError, KeyError) as e:
                print("אזהרת פתיחת הפטרה", f"לא ניתן למשוך פתיחת הפטרה מ-Sefaria.\n{e}")

        return ""

    @classmethod
    def extract_haftara_for_date(cls, items, date_str, include_opening_lookup=False):
        for item in items:
            if item.get("date", "")[:10] != date_str:
                continue

            haft = item.get("haft", {})
            if haft.get("k") and haft.get("b") and haft.get("e"):
                haftara = cls.format_hebrew_tanach_ref(haft["k"], haft["b"], haft["e"])
                opening = cls.extract_haftara_opening_for_item(item, include_opening_lookup)
                if opening:
                    return f"{haftara} - {opening}"
                return haftara

            if item.get("haftara"):
                return item["haftara"]

            leyning = item.get("leyning", {})
            if leyning.get("haftarah"):
                return leyning["haftarah"]

        return ""

    def get_upcoming_saturday(self):
        today = datetime.now()
        days_ahead = 5 - today.weekday()
        if days_ahead < 0:
            days_ahead += 7
        return (today + timedelta(days=days_ahead)).date()

    def load_parashiyot(self):
        try:
            current_year = datetime.now().year
            url = "https://www.hebcal.com/hebcal"
            response = requests.get(url, params=self.build_parasha_list_params(current_year))
            response.raise_for_status()
            data = response.json()

            today = datetime.now().date()
            upcoming_saturday = self.get_upcoming_saturday()

            for item in data.get("items", []):
                if item.get("category") == "parashat":
                    date_str = item.get("date")
                    parasha_date = datetime.strptime(date_str, "%Y-%m-%d").date()

                    if parasha_date >= today - timedelta(days=3):
                        heb_name = item.get("hebrew", item.get("title"))
                        self.parashiyot[heb_name] = date_str

            parasha_list = list(self.parashiyot.keys())
            self.parasha_dropdown.addItems(parasha_list)

            if parasha_list:
                default_index = 0
                for idx, name in enumerate(parasha_list):
                    if self.parashiyot[name] == upcoming_saturday.strftime("%Y-%m-%d"):
                        default_index = idx
                        break
                self.parasha_dropdown.setCurrentIndex(default_index)

        except Exception as e:
            QMessageBox.critical(self, "שגיאת רשת", f"לא ניתן למשוך רשימת פרשות.\n{e}")

    def fetch_chai_tables_netz(self, hebrew_year, metro_area):
        """Scrapes the Chai Tables HTML and builds a dictionary mapping month and day to Netz times."""
        url = "https://chaitables.com/cgi-bin/ChaiTables.cgi/"
        payload = {
            "cgi_TableType": "BY",
            "cgi_country": "Eretz Yisroel",
            "cgi_USAcities1": "1",
            "cgi_USAcities2": "0",
            "cgi_searchradius": "",
            "cgi_Placename": "?",
            "cgi_eroslatitude": "0.0",
            "cgi_eroslongitude": "0.0",
            "cgi_eroshgt": "0.0",
            "cgi_geotz": "2",
            "cgi_exactcoord": "OFF",
            "cgi_MetroArea": metro_area,
            "cgi_types": "0",
            "cgi_RoundSecond": "-1",
            "cgi_AddCushion": "0",
            "cgi_24hr": "ON",
            "cgi_DST": "ON",
            "cgi_typezman": "-1",
            "cgi_yrheb": str(hebrew_year),
            "cgi_optionheb": "0",
            "cgi_UserNumber": "49939",
            "cgi_Language": "Hebrew",
            "cgi_AllowShaving": "OFF"
        }
        try:
            response = requests.get(url, params=payload)

            response.encoding = 'utf-8'

            response.raise_for_status()
            soup = BeautifulSoup(response.text, 'html.parser')

            tables = soup.find_all('table')
            data_table = None
            for tbl in tables:
                if 'תשרי' in tbl.get_text():
                    data_table = tbl
                    break

            if not data_table:
                print("Could not find the Chai Tables data grid in the HTML.")
                return {}

            rows = data_table.find_all('tr')

            # Map Hebrew headers in Chai Tables to Hebcal's internal English mapping
            hebcal_month_map = {
                "תשרי": "Tishrei", "מרחשון": "Cheshvan", "כסלו": "Kislev",
                "טבת": "Tevet", "שבט": "Shvat", "אדר": "Adar",
                "אדר א": "Adar I", "אדר ב": "Adar II",
                "ניסן": "Nisan", "אייר": "Iyyar", "סיון": "Sivan",
                "תמוז": "Tamuz", "אב": "Av", "אלול": "Elul"
            }

            col_to_month = {}
            start_row = 0

            # Dynamically hunt for the header row
            for r_idx, row in enumerate(rows):
                cells = row.find_all(['td', 'th'])
                for c_idx, cell in enumerate(cells):
                    text_clean = cell.get_text(strip=True).replace('"', '').replace("'", "")
                    for heb, eng in hebcal_month_map.items():
                        if heb in text_clean:
                            col_to_month[c_idx] = eng
                            break
                # Once we found the row with the months, mark our starting point and stop hunting
                if col_to_month:
                    print(f"Successfully mapped columns: {col_to_month}")
                    start_row = r_idx + 1
                    break

            netz_data = {month: {} for month in col_to_month.values()}

            # Iterate through the actual days
            for day_idx in range(start_row, len(rows)):
                cells = rows[day_idx].find_all('td')
                if len(cells) < 3: continue

                # The first cell in the data row is the Hebrew day number (1 to 30)
                try:
                    # Clean the day text (e.g. converting 'א' to 1 is tricky, so we rely on the loop index offset)
                    hebrew_day = day_idx - start_row + 1
                except Exception:
                    continue

                for col_idx, month_eng in col_to_month.items():
                    if col_idx < len(cells):
                        # get_text automatically cleans out HTML tags like <u> and <b>
                        time_str = cells[col_idx].get_text(strip=True)
                        if time_str and ":" in time_str:
                            parts = time_str.split(':')
                            if len(parts) == 3:
                                h = parts[0].zfill(2)
                                netz_data[month_eng][hebrew_day] = f"{h}:{parts[1]}:{parts[2]}"

            print(f"Successfully cached Chai Tables for {hebrew_year}")
            return netz_data
        except Exception as e:
            print("אזהרת צ'יי", f"שגיאה במשיכת לוחות צ'יי, נשתמש בנתוני גיבוי מ-Hebcal.\n{e}")
            return {}

    def fetch_week_zmanim(self, hebcal_id, chai_metro, start_date_str):
        start_date = datetime.strptime(start_date_str, "%Y-%m-%d")

        zmanim_mapping = {
            "alotHaShachar": "עלות השחר",
            "misheyakir": "זמן טלית ותפילין",
            "sunrise": "הנץ החמה (הנראה)"
        }

        week_data = {hebrew_name: [] for hebrew_name in zmanim_mapping.values()}
        hebrew_dates = []
        daf_yomi_week = []
        special_events_week = ["--"] * 7
        yom_tov_week = [False] * 7
        haftara = ""

        try:
            self.generate_btn.setText("מושך נתוני דף יומי וימים מיוחדים...")
            QApplication.processEvents()
            end_date_str = (start_date + timedelta(days=6)).strftime("%Y-%m-%d")
            hebcal_week_url = "https://www.hebcal.com/hebcal"
            hebcal_week_params = self.build_hebcal_week_params(hebcal_id, start_date_str, end_date_str)

            daf_dict = {}
            hebcal_week_response = requests.get(hebcal_week_url, params=hebcal_week_params)
            if hebcal_week_response.status_code == 200:
                hebcal_week_items = hebcal_week_response.json().get("items", [])
                special_events_week = self.build_special_events_week(start_date, hebcal_week_items)
                yom_tov_week = self.build_yom_tov_week(start_date, hebcal_week_items)
                for item in hebcal_week_items:
                    if item.get("category") == "dafyomi":
                        daf_dict[item.get("date")] = item.get("hebrew", "")

            try:
                leyning_response = requests.get(
                    "https://www.hebcal.com/leyning",
                    params={
                        "cfg": "json",
                        "date": start_date_str,
                        "i": "on",
                        "triennial": "off",
                    },
                )
                if leyning_response.status_code == 200:
                    haftara = self.extract_haftara_for_date(
                        leyning_response.json().get("items", []),
                        start_date_str,
                        include_opening_lookup=True,
                    )
            except requests.RequestException as e:
                print("אזהרת הפטרה", f"לא ניתן למשוך הפטרה מ-Hebcal.\n{e}")

            for i in range(7):
                self.generate_btn.setText(f"מושך נתונים... (יום {i + 1}/7)")
                QApplication.processEvents()

                current_date = start_date + timedelta(days=i)
                date_str = current_date.strftime("%Y-%m-%d")

                daf_yomi_week.append(daf_dict.get(date_str, "--"))

                # 1. Fetch Hebcal Hebrew Date Converter
                converter_url = f"https://www.hebcal.com/converter?cfg=json&date={date_str}&g2h=1&strict=1"
                converter_response = requests.get(converter_url)
                converter_response.raise_for_status()
                converter_data = converter_response.json()

                hebrew_date_str = converter_data.get("hebrew", "")
                hebrew_dates.append(hebrew_date_str)

                # Extract year, month, and day for Chai Tables mapping
                hy = converter_data.get("hy")
                hm = converter_data.get("hm")
                hd = converter_data.get("hd")

                # 2. Fetch/Cache Chai Tables (Only triggered once per year)
                if hy not in self.chai_tables_cache:
                    self.generate_btn.setText(f"מושך לוחות צ'יי לשנת {hy}...")
                    QApplication.processEvents()
                    self.chai_tables_cache[hy] = self.fetch_chai_tables_netz(hy, chai_metro)

                # Retrieve the specific Netz time for today from the Chai Tables Cache
                chai_netz = None
                if hy in self.chai_tables_cache and hm in self.chai_tables_cache[hy]:
                    chai_netz = self.chai_tables_cache[hy][hm].get(hd)

                # 3. Fetch standard Hebcal Zmanim
                zmanim_url = f"https://www.hebcal.com/zmanim?cfg=json&geonameid={hebcal_id}&date={date_str}&sec=1"
                zmanim_response = requests.get(zmanim_url)
                zmanim_response.raise_for_status()
                zmanim_data = zmanim_response.json()
                times = zmanim_data.get("times", {})

                for key, hebrew_name in zmanim_mapping.items():
                    # HYBRID INJECTION: If we are looking for Netz AND Chai Tables has a time for today, inject it!
                    if hebrew_name == "הנץ החמה (הנראה)" and chai_netz:
                        week_data[hebrew_name].append(chai_netz)
                    # Otherwise, use standard Hebcal parsed time
                    elif key in times:
                        time_iso = times[key]
                        time_str = time_iso.split('+')[0]
                        dt_obj = datetime.strptime(time_str, "%Y-%m-%dT%H:%M:%S")
                        week_data[hebrew_name].append(dt_obj.strftime("%H:%M:%S"))
                    else:
                        week_data[hebrew_name].append("--:--:--")

            return start_date, week_data, zmanim_mapping, hebrew_dates, daf_yomi_week, special_events_week, haftara, yom_tov_week

        except Exception as e:
            QMessageBox.critical(self, "שגיאת רשת", f"שגיאה במשיכת הנתונים: {e}")
            return None, None, None, None, None, None, None, None

    def open_file_in_os(self, filename):
        try:
            if sys.platform == "win32":
                os.startfile(filename)
            elif sys.platform == "darwin":
                subprocess.call(["open", filename])
            else:
                subprocess.call(["xdg-open", filename])
        except Exception as e:
            QMessageBox.warning(self, "אזהרה", f"הקובץ נוצר אך לא ניתן היה לפתוח אותו אוטומטית.\n{e}")

    def generate_schedule(self):
        city_name = self.city_dropdown.currentText()
        hebcal_id, chai_metro = self.cities[city_name]

        parasha_name = self.parasha_dropdown.currentText()
        if not parasha_name:
            QMessageBox.warning(self, "אזהרה", "אנא בחר פרשת שבוע תקינה.")
            return

        start_date_str = self.parashiyot[parasha_name]
        comments = self.comments_text.toPlainText().strip()

        self.generate_btn.setEnabled(False)

        start_date, week_data, zmanim_mapping, hebrew_dates, daf_yomi_week, special_events_week, haftara, yom_tov_week = self.fetch_week_zmanim(
            hebcal_id, chai_metro, start_date_str)

        if week_data and hebrew_dates:
            try:
                filename = self.export_to_excel(city_name, parasha_name, start_date, week_data, comments,
                                                zmanim_mapping, hebrew_dates, daf_yomi_week, special_events_week,
                                                haftara, yom_tov_week)
                self.open_file_in_os(filename)
            except PermissionError:
                QMessageBox.critical(self, "שגיאת שמירה", "הקובץ כבר פתוח ב-Excel.\nאנא סגור אותו ונסה שוב.")
            except Exception as e:
                QMessageBox.critical(self, "שגיאה", f"שגיאה בשמירת הקובץ: {e}")

        self.generate_btn.setText("עדכן והכן להדפסה")
        self.generate_btn.setEnabled(True)

    def export_to_excel(self, city, parasha_name, start_date, week_data, comments, zmanim_mapping, hebrew_dates,
                        daf_yomi_week, special_events_week=None, haftara="", yom_tov_week=None):
        date_str = start_date.strftime("%Y-%m-%d")
        safe_parasha_name = "".join(c for c in parasha_name if c.isalnum() or c in " -")
        filename = f"לוח_זמנים_שבועי_{safe_parasha_name}_{date_str}.xlsx"
        clean_parasha = parasha_name.replace("פרשת ", "").strip()
        hebrew_year, daily_hebrew_dates = ZmanimApp.split_hebrew_year_from_dates(hebrew_dates)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "לוח זמנים"
        ws.sheet_view.rightToLeft = True
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.25
        ws.page_margins.bottom = 0.25
        ws.page_margins.header = 0.1
        ws.page_margins.footer = 0.1
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.print_options.horizontalCentered = True

        default_font = Font(name="Arial", size=14)
        title_font = Font(name="Arial", size=20, bold=True)
        subtitle_font = Font(name="Arial", size=16, bold=True)
        schedule_title_font = Font(name="David", size=16, bold=True, underline="single")
        hebrew_date_font = Font(name="David", size=14, italic=True)
        day_header_font = Font(name="David", size=14, bold=True, underline="single")
        bold_font = Font(name="Arial", size=14, bold=True)
        italic_font = Font(name="Arial", size=14, italic=True)
        center_alignment = Alignment(horizontal="center")
        fit_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, shrink_to_fit=True)
        wrapped_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, shrink_to_fit=False)

        ws.append(["בס\"ד", "", "", "", "", "", "", "", date_str])
        ws.append(["", hebrew_year])
        ws.merge_cells("B2:I2")
        ws.cell(row=2, column=2).font = subtitle_font
        ws.cell(row=2, column=2).alignment = center_alignment

        ws.append(["", "בית הכנסת משמרת  --  ותיקין"])
        ws.row_dimensions[3].height = 26.25
        ws.merge_cells("B3:I3")
        ws.cell(row=3, column=2).font = title_font
        ws.cell(row=3, column=2).alignment = center_alignment
        ws.append([])

        schedule_title = f'לוח זמנים לשבת פרשת "{clean_parasha}"'
        if haftara:
            schedule_title = f"{schedule_title} | הפטרה: {haftara}"

        ws.append(["", schedule_title])
        ws.row_dimensions[5].height = 20.25
        ws.merge_cells("B5:I5")
        ws.cell(row=5, column=2).font = schedule_title_font
        ws.cell(row=5, column=2).alignment = center_alignment
        ws.append([])

        ws.append(["", ""] + daily_hebrew_dates)
        ws.row_dimensions[7].height = 15
        for col_idx in range(3, 10):
            cell = ws.cell(row=7, column=col_idx)
            cell.font = hebrew_date_font
            cell.alignment = center_alignment

        ws.append(["", "", "שבת", "א'", "ב'", "ג'", "ד'", "ה'", "ו'"])
        ws.row_dimensions[8].height = 18
        for col_idx in range(3, 10):
            cell = ws.cell(row=8, column=col_idx)
            cell.font = day_header_font
            cell.alignment = center_alignment

        ws.append([])

        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 3
        for col in range(3, 10):
            ws.column_dimensions[get_column_letter(col)].width = 17

        current_row = 10
        for hebrew_name in zmanim_mapping.values():
            row_data = [hebrew_name, ""] + week_data[hebrew_name]
            ws.append(row_data)
            ws.row_dimensions[current_row].height = 36

            for col_idx in range(1, 10):
                ws.cell(row=current_row, column=col_idx).font = default_font
                ws.cell(row=current_row, column=col_idx).alignment = fit_alignment

            for col_idx in range(3, 10):
                ws.cell(row=current_row, column=col_idx).alignment = fit_alignment

            ws.append([])
            ws.row_dimensions[current_row + 1].height = 18
            for col_idx in range(1, 10):
                ws.cell(row=current_row + 1, column=col_idx).font = default_font
            current_row += 2

        vatikin_times = []
        netz_times = week_data.get("הנץ החמה (הנראה)", ["--:--:--"] * 7)
        if yom_tov_week is None:
            yom_tov_week = [False] * 7

        for i, netz_str in enumerate(netz_times):
            if netz_str == "--:--:--":
                vatikin_times.append("--:--:--")
            else:
                netz_dt = datetime.strptime(netz_str, "%H:%M:%S")
                offset_minutes = 32 if i == 0 or yom_tov_week[i] else 22
                vatikin_dt = netz_dt - timedelta(minutes=offset_minutes)
                vatikin_times.append(vatikin_dt.strftime("%H:%M:%S"))

        ws.append(["שחרית ותיקין", ""] + vatikin_times)
        ws.row_dimensions[current_row].height = 30
        for col_idx in range(1, 10):
            ws.cell(row=current_row, column=col_idx).font = default_font
            ws.cell(row=current_row, column=col_idx).alignment = fit_alignment
        for col_idx in range(3, 10):
            ws.cell(row=current_row, column=col_idx).alignment = fit_alignment
            ws.cell(row=current_row, column=col_idx).font = bold_font

        ws.append([])
        ws.row_dimensions[current_row + 1].height = 18
        for col_idx in range(1, 10):
            ws.cell(row=current_row + 1, column=col_idx).font = default_font
        current_row += 2

        # Add the Daf Yomi row
        ws.append(["דף יומי", ""] + daf_yomi_week)
        ws.row_dimensions[current_row].height = 24
        ws.cell(row=current_row, column=1).font = bold_font
        ws.cell(row=current_row, column=2).font = default_font
        ws.cell(row=current_row, column=1).alignment = wrapped_alignment
        ws.cell(row=current_row, column=2).alignment = wrapped_alignment
        for col_idx in range(3, 10):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.alignment = wrapped_alignment
            cell.font = italic_font

        ws.append([])
        ws.row_dimensions[current_row + 1].height = 18
        for col_idx in range(1, 10):
            ws.cell(row=current_row + 1, column=col_idx).font = default_font
        current_row += 2

        if special_events_week is None:
            special_events_week = ["--"] * 7

        ws.append(["ימים מיוחדים", ""] + special_events_week)
        ws.row_dimensions[current_row].height = ZmanimApp.calculate_content_row_height(special_events_week)
        ws.cell(row=current_row, column=1).font = bold_font
        ws.cell(row=current_row, column=2).font = default_font
        ws.cell(row=current_row, column=1).alignment = wrapped_alignment
        ws.cell(row=current_row, column=2).alignment = wrapped_alignment
        for col_idx in range(3, 10):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.alignment = wrapped_alignment
            cell.font = italic_font

        ws.append([])
        ws.row_dimensions[current_row + 1].height = 18
        for col_idx in range(1, 10):
            ws.cell(row=current_row + 1, column=col_idx).font = default_font

        ws.append(["", "הערות:"])
        ws.row_dimensions[ws.max_row].height = 18
        ws.cell(row=ws.max_row, column=2).font = bold_font

        if comments:
            for line in comments.split('\n'):
                ws.append(["", "", line])
                ws.row_dimensions[ws.max_row].height = 18
                ws.cell(row=ws.max_row, column=3).font = default_font

        ws.append([])
        ws.append(["", "שבת שלום ומבורך בשורות טובות"])
        ws.row_dimensions[ws.max_row].height = 18
        ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=9)
        ws.cell(row=ws.max_row, column=2).font = bold_font
        ws.cell(row=ws.max_row, column=2).alignment = center_alignment

        ws.print_area = f"A1:I{ws.max_row}"

        wb.save(filename)
        return filename


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = ZmanimApp()
    window.show()
    sys.exit(app.exec())
